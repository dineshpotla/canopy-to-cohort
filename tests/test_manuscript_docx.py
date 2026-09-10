"""Read-only structural checks for the generated coauthor-review manuscript.

Run with bundled Python after make paper-word. These tests do not replace
the required page-image inspection and do not approve scientific claims.
"""
import csv
import json
from pathlib import Path
import unittest
from zipfile import ZipFile

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "outputs/paper/michigan-recruitment-paper.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@unittest.skipUnless(PATH.is_file(), "Build the Word manuscript first")
class ManuscriptChecks(unittest.TestCase):
    def setUp(self):
        self.doc = Document(PATH)
        self.paragraphs = [p.text for p in self.doc.paragraphs]

    def section_text(self, start, end):
        return " ".join(self.paragraphs[self.paragraphs.index(start) + 1:self.paragraphs.index(end)])

    def test_candidate_length_limits(self):
        self.assertLessEqual(len(self.paragraphs[0].split()), 15)
        self.assertLessEqual(len(self.section_text("Abstract", "Study implications").split()), 200)
        self.assertLessEqual(len(self.section_text("Study implications", "Introduction").split()), 100)
        self.assertEqual(self.doc.core_properties.title, self.paragraphs[0])

    def test_tables_figures_and_alt_text(self):
        self.assertEqual(len(self.doc.tables), 16)
        self.assertEqual(len(self.doc.inline_shapes), 2)
        for table in self.doc.tables:
            self.assertTrue(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
        for shape in self.doc.inline_shapes:
            self.assertTrue(shape._inline.docPr.get("descr"))

    def test_equations_remain_editable_word_math(self):
        with ZipFile(PATH) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            equations = root.xpath("//m:oMath", namespaces={
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"})
            self.assertGreaterEqual(len(equations), 8)

    def table_rows(self, required_headers):
        matches = [t for t in self.doc.tables
                   if set(required_headers) <= {c.text for c in t.rows[0].cells}]
        self.assertEqual(len(matches), 1, required_headers)
        table = matches[0]
        headers = [c.text for c in table.rows[0].cells]
        return [dict(zip(headers, [c.text for c in row.cells])) for row in table.rows[1:]]

    def source_rows(self, name):
        with (ROOT / "outputs/tables" / f"{name}.csv").open() as source:
            return list(csv.DictReader(source))

    def test_analytical_detail_is_in_the_document(self):
        for heading in ("Analytical appendix", "County fold support and fitting diagnostics",
                        "Absolute performance uncertainty", "Paired ranking and log loss contrasts",
                        "Full population penalized coefficients", "Joint count and sapling strata"):
            self.assertIn(heading, self.paragraphs)
        text = " ".join(self.paragraphs)
        for detail in ("No predictor imputation is used", "frequency weight",
                       "no calibration-parameter uncertainty intervals",
                       "whose interval excludes zero even though the Brier interval crosses zero"):
            self.assertIn(detail, text)

    def test_coefficient_table_matches_existing_exports(self):
        actual = self.table_rows(["Model", "Term", "All", "Detected"])
        self.assertEqual(len(actual), 14)
        model_names = {"Prevalence": "Prevalence", "Observation design": "Design",
                       "Seedling count and design": "Count + design",
                       "Maple stages and design": "Stages + design"}
        term_names = {"(Intercept)": "Intercept", "z_coverage": "Coverage", "z_interval": "Interval",
                      "z_seedling_count": "Log seedling count", "z_sapling_tpa": "Log sapling density",
                      "z_established_ba": "Log established basal area"}
        lookup = {(r["Model"], r["Term"]): r for r in actual}
        for row in self.source_rows("recruitment-model-coefficients"):
            column = "All" if row["population"] == "All eligible pairs" else "Detected"
            value = lookup[(model_names[row["model"]], term_names[row["term"]])][column]
            self.assertEqual(value, f'{float(row["penalized_log_odds_coefficient"]):.4f}')

    def test_sensitivity_table_preserves_matched_denominators(self):
        actual = self.table_rows(["Restriction", "N / events", "Brier count / stage", "Difference"])
        self.assertEqual(len(actual), 8)
        lookup = {(r["Population"], r["Restriction"]): r for r in actual}
        labels = {"Exact mapped overlap": "Exact overlap", "Follow-up manual >=9": "Manual at least 9",
                  "Ridge lambda 4": "Penalty 4", "Unpenalized": "Penalty 0"}
        source = self.source_rows("recruitment-model-sensitivities")
        for row in source:
            if row["model"] != "Seedling count and design":
                continue
            stage = next(r for r in source if r["model"] == "Maple stages and design"
                         and r["population"] == row["population"] and r["sensitivity"] == row["sensitivity"])
            self.assertEqual((row["observations"], row["events"]), (stage["observations"], stage["events"]))
            pop = "All" if row["population"] == "All eligible pairs" else "Detected"
            rendered = lookup[(pop, labels[row["sensitivity"]])]
            self.assertEqual(rendered["N / events"], f'{row["observations"]} / {row["events"]}')
            self.assertEqual(rendered["Brier count / stage"],
                             f'{float(row["brier_score"]):.5f} / {float(stage["brier_score"]):.5f}')
            self.assertEqual(rendered["Difference"], f'{float(stage["brier_score"])-float(row["brier_score"]):.5f}')

    def test_extended_contrasts_and_all_survey_policies_are_present(self):
        contrasts = self.table_rows(["Addition", "Metric", "Difference", "95% interval"])
        self.assertEqual(len(contrasts), 12)
        lookup = {(r["Population"], r["Addition"], r["Metric"]): r for r in contrasts}
        for row in self.source_rows("recruitment-model-paired-differences"):
            metrics = {"log_loss": "Log loss", "roc_auc": "AUC", "average_precision": "AP"}
            if row["metric"] not in metrics:
                continue
            key = ("All" if row["population"] == "All eligible pairs" else "Detected",
                   "Count to design" if row["model"] == "Seedling count and design" else "Stages to count",
                   metrics[row["metric"]])
            self.assertEqual(lookup[key]["Difference"], f'{float(row["difference"]):.4f}')
            self.assertEqual(lookup[key]["95% interval"], f'{float(row["lower"]):.4f} to {float(row["upper"]):.4f}')
        policies = self.table_rows(["Strategy", "Expected targets", "All targets captured"])
        self.assertEqual(len(policies), 6)
        lookup = {r["Strategy"]: r for r in policies}
        for row in self.source_rows("recruitment-survey-summary"):
            if (row["population"] == "Baseline seedlings detected" and row["objective"] == "No recorded entry"
                    and float(row["requested_fraction"]) == 0.25):
                self.assertEqual(lookup[row["policy"]]["Expected targets"], f'{float(row["expected_targets"]):.2f}')

    def test_review_comments_are_anchored(self):
        expected = json.loads((ROOT / "research/runs/recruitment-2026-09-08/manuscript-review-comments.json").read_text())
        self.assertEqual(len(self.doc.comments), len(expected))
        with ZipFile(PATH) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            comments = etree.fromstring(archive.read("word/comments.xml"))
            def ids(kind):
                return root.xpath(f"//w:{kind}/@w:id", namespaces=NS)
            actual = comments.xpath("//w:comment/@w:id", namespaces=NS)
            self.assertCountEqual(actual, ids("commentRangeStart"))
            self.assertCountEqual(actual, ids("commentRangeEnd"))
            self.assertCountEqual(actual, ids("commentReference"))
            for item in expected:
                self.assertTrue(any(item["anchor"] in p for p in self.paragraphs))

    def test_availability_does_not_claim_a_release(self):
        text = self.section_text("Data and code availability", "Computational workflow")
        self.assertIn("has not yet been assigned", text)
        self.assertIn("2c1eb908a47436d4edd0f3ba9e3d647b79a4f36cef972c5ba83277300817aee1", text)
        self.assertIn("Use of AI assistance", self.paragraphs)

    def test_no_repository_only_word_hyperlinks(self):
        for link in self.doc._element.xpath(".//w:hyperlink"):
            rid = link.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if rid:
                target = str(self.doc.part.rels[rid].target_ref)
                self.assertTrue(target.startswith(("https://", "http://", "mailto:")), target)
        text = " ".join(self.paragraphs)
        for marker in ("`r ", "TODO", "TBD", "turn125view", "cite"):
            self.assertNotIn(marker, text)


if __name__ == "__main__":
    unittest.main()
