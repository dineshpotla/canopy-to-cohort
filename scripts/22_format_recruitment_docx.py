"""Apply plain academic formatting to the Quarto-generated editable manuscript.

Run with the bundled workspace Python. Content and computed values remain in
report/recruitment.qmd; this script only normalizes the Word presentation.
"""
import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def element(tag, **attributes):
    result = OxmlElement(tag)
    for name, value in attributes.items():
        result.set(qn("w:" + name), str(value))
    return result


def format_manuscript(source, target, review_comments=None, manuscript_source=None):
    doc = Document(source)
    # Quarto wraps a caption and its table/figure in a one-cell layout table.
    # Unwrap these containers so only actual data tables receive table styling.
    for outer in list(doc.tables):
        if len(outer.rows) == 1 and len(outer.columns) == 1 and (
            outer.cell(0, 0).tables or outer._tbl.xpath(".//w:drawing")
        ):
            parent = outer._tbl.getparent()
            position = parent.index(outer._tbl)
            for child in list(outer.cell(0, 0)._tc):
                if child.tag in [qn("w:p"), qn("w:tbl")]:
                    parent.insert(position, child)
                    position += 1
            parent.remove(outer._tbl)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.85)
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.35)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Page ")
        footer._p.append(element("w:fldSimple", instr="PAGE"))
        for run in footer.runs:
            run.font.size = Pt(10)
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.underline = False
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 19), ("Subtitle", 11), ("Heading 1", 15),
                       ("Heading 2", 13), ("Heading 3", 11.5)]:
        if name in doc.styles:
            style = doc.styles[name]
            style.font.size = Pt(size)
            style.font.bold = name != "Subtitle"
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.page_break_before = False
    for name in ["Caption", "Image Caption", "Table Caption"]:
        if name in doc.styles:
            doc.styles[name].font.size = Pt(10)
            doc.styles[name].paragraph_format.space_after = Pt(5)
    for name in ["Source Code", "Verbatim Char"]:
        if name in doc.styles:
            doc.styles[name].font.name = "Courier New"
            doc.styles[name].font.size = Pt(9)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        label = paragraph.text.replace("\u00a0", " ")
        if label.startswith("Table "):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if label in ["Analytical appendix", "References"]:
            paragraph.paragraph_format.page_break_before = True
        if paragraph.style.name in ["Title", "Subtitle"]:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Keep an image and its following caption together.
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style.name == "Table Caption":
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style.name in ["Bibliography", "References"]:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_after = Pt(7)
        if paragraph.style.name == "Title":
            ppr = paragraph._p.get_or_add_pPr()
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    for table in doc.tables:
        table.autofit = False
        columns = len(table.columns)
        headers = [cell.text for cell in table.rows[0].cells]
        if "Source quantity and unit" in headers:
            widths = [1.35, 2.7, 1.5, 0.95]
        elif "Brier count / stage" in headers:
            widths = [0.9, 1.1, 0.75, 1.6, 0.75, 1.4]
        elif "Brier estimate and interval" in headers:
            widths = [0.9, 1.4, 2.2, 2.0]
        elif "Joint intercept" in headers:
            widths = [0.8, 1.6, 1.15, 1.25, 0.85, 0.85]
        elif "Train N / events" in headers:
            widths = [1.05, 0.5, 1.85, 1.85, 1.25]
        elif "Expected targets" in headers:
            widths = [1.85, 0.95, 1.2, 1.4, 1.1]
        elif "Term" in headers and "Detected" in headers:
            widths = [1.6, 2.4, 1.25, 1.25]
        elif "Metric" in headers and "Difference" in headers:
            widths = [0.9, 1.3, 0.75, 0.95, 2.6]
        elif "Existing saplings" in headers:
            widths = [1.15, 1.55, 1.3, 1.2, 1.3]
        elif "Share of comparable stems" in headers:
            widths = [3.85, 0.85, 1.8]
        elif "With new tally" in headers and columns == 5:
            widths = [1.1, 1.1, 1.15, 1.15, 2.0]
        elif columns == 2:
            widths = [5.35, 1.15]
        elif columns == 4 and "Difference" in headers:
            widths = [1.45, 1.35, 0.85, 2.85]
        elif columns == 4:
            widths = [1.8, 1.6, 1.55, 1.55]
        elif columns == 6 and "Calibration slope" in headers:
            widths = [1.75, 0.95, 0.9, 0.8, 0.8, 1.3]
        elif columns == 6 and "Test N" in headers:
            widths = [1.45, 1.65, 0.7, 0.7, 1.05, 0.95]
        elif columns == 6:
            widths = [1.55, 1.05, 0.85, 0.9, 1.0, 1.15]
        else:
            widths = [6.5 / columns] * columns
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        props = table._tbl.tblPr
        for tag in ["w:tblStyle", "w:tblBorders", "w:tblCellMar", "w:tblW"]:
            for old in list(props.findall(qn(tag))):
                props.remove(old)
        props.append(element("w:tblW", w=9360, type="dxa"))
        borders = element("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            borders.append(element("w:" + edge, val="single", sz=4, color="D9D9D9"))
        props.append(borders)
        margins = element("w:tblCellMar")
        for edge, value in [("top", 75), ("bottom", 75), ("left", 90), ("right", 90)]:
            margins.append(element("w:" + edge, w=value, type="dxa"))
        props.append(margins)
        for index, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(element("w:cantSplit"))
            if index == 0:
                for old in list(trpr.findall(qn("w:tblHeader"))):
                    trpr.remove(old)
                trpr.append(element("w:tblHeader", val="true"))
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                # Remove Pandoc cell-level black rules so the specified light
                # table borders apply uniformly, including below the header.
                for border in list(cell._tc.get_or_add_tcPr().findall(qn("w:tcBorders"))):
                    cell._tc.get_or_add_tcPr().remove(border)
                if index == 0:
                    cell._tc.get_or_add_tcPr().append(element("w:shd", fill="F2F2F2", val="clear"))
                for paragraph in cell.paragraphs:
                    alignments = paragraph._p.get_or_add_pPr().findall(qn("w:jc"))
                    for old in alignments[:-1]:
                        paragraph._p.get_or_add_pPr().remove(old)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.05
                    # Long analytical tables may continue across pages with
                    # repeated headers; chaining all rows creates large gaps.
                    paragraph.paragraph_format.keep_with_next = (
                        index < len(table.rows) - 1 if len(table.rows) <= 7 else index == 0
                    )
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.bold = index == 0
                        run.font.color.rgb = RGBColor(0, 0, 0)
    # Repository-relative links do not travel with a standalone Word file.
    # Preserve their labels; retain DOI/web links and internal citation anchors.
    for link in list(doc._element.xpath(".//w:hyperlink")):
        rid = link.get(qn("r:id"))
        if rid and rid in doc.part.rels:
            destination = doc.part.rels[rid].target_ref
            if not str(destination).startswith(("https://", "http://", "mailto:")):
                parent = link.getparent()
                position = parent.index(link)
                for child in list(link):
                    parent.insert(position, child)
                    position += 1
                parent.remove(link)
    titles = [p.text for p in doc.paragraphs if p.style.name == "Title"]
    if len(titles) != 1:
        raise ValueError("Expected exactly one manuscript title")
    doc.core_properties.title = titles[0]
    doc.core_properties.subject = "Michigan FIA core research manuscript"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "sugar maple; FIA; regeneration; sapling recruitment; Michigan"
    if manuscript_source:
        # Quarto's Word export can leave docPr descr empty despite fig-alt.
        # Read the canonical descriptions, in figure order, instead of allowing
        # a second hand-maintained set of accessibility text to drift.
        descriptions = [json.loads(value) for value in re.findall(
            r'^#\| fig-alt:\s*(.+)$', manuscript_source.read_text(), flags=re.MULTILINE)]
        if len(descriptions) != len(doc.inline_shapes):
            raise ValueError("Each manuscript figure requires a canonical alternative description")
        for shape, description in zip(doc.inline_shapes, descriptions):
            if not isinstance(description, str) or not description.strip():
                raise ValueError("Figure alternative description must be nonempty text")
            shape._inline.docPr.set("descr", description)
    if review_comments:
        for item in json.loads(review_comments.read_text()):
            matches = [p for p in doc.paragraphs if item["anchor"] in p.text]
            if len(matches) != 1 or not matches[0].runs:
                raise ValueError(f"Review comment must match exactly one paragraph: {item['anchor']}")
            doc.add_comment(matches[0].runs, text=item["text"], author="Codex", initials="CX")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(f"Saved {target}; {len(doc.tables)} editable tables; {len(doc.inline_shapes)} figures")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    parser.add_argument("--review-comments", type=Path)
    parser.add_argument("--manuscript-source", type=Path)
    args = parser.parse_args()
    format_manuscript(args.source, args.target, args.review_comments, args.manuscript_source)
